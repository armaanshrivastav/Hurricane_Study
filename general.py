import re
import pandas as pd
from openpyxl import load_workbook

def append_to_excel(file_name, tab_name, text, correlation_value):
    """
    Appends the provided text and correlation value to the next available line
    in a specified tab of an Excel file.

    Parameters:
    - file_name (str): The name of the Excel file (including file extension).
    - tab_name (str): The name of the tab to create or modify.
    - text (str): The text to append to the tab.
    - correlation_value (float): The correlation value to append to the tab.
    """

    print(tab_name)
    # Create a DataFrame for the new row
    new_data = pd.DataFrame({
        "Description": [text],
        "Correlation Value": [correlation_value]
    })

    try:
        # Load the existing workbook
        book = load_workbook(file_name)
        
        if tab_name in book.sheetnames:
            # Load existing data from the specified tab
            existing_data = pd.read_excel(file_name, sheet_name=tab_name)
            # Append the new data to the existing data
            updated_data = pd.concat([existing_data, new_data], ignore_index=True)
        else:
            # If the tab doesn't exist, use the new data
            updated_data = new_data

        # Write the updated data back to the file
        with pd.ExcelWriter(file_name, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
            updated_data.to_excel(writer, sheet_name=tab_name, index=False)

    except FileNotFoundError:
        # If the file doesn't exist, create it and write the new data
        with pd.ExcelWriter(file_name, engine='openpyxl') as writer:
            new_data.to_excel(writer, sheet_name=tab_name, index=False)

    print(f"Data appended to tab '{tab_name}' in '{file_name}'.")


# Function to strip numeric characters at the end of a string
def strip_numeric_suffix(input_string):
    return re.sub(r'\d+$', '', input_string)

print(strip_numeric_suffix('Helene1'))
